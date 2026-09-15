#!/usr/bin/env python3
"""Generate a single-page architecture PPT for asc-devkit / asc-comm."""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT_DIR = Path(__file__).resolve().parent
PPTX_PATH = OUT_DIR / "asc_devkit_comm_arch.pptx"
PNG_PATH = OUT_DIR / "asc_devkit_comm_arch.png"

SLIDE_W = 13.333
SLIDE_H = 7.5
PNG_W = 1920
PNG_H = 1080

FONT_PATH = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"
EA_FONT = "Microsoft YaHei"
LATIN_FONT = "Calibri"

# Palette
BG = (244, 247, 251)
NAVY = (15, 39, 68)
NAVY_DEEP = (10, 28, 50)
RED = (207, 10, 44)
WHITE = (255, 255, 255)
MUTED = (90, 107, 128)
LINE = (210, 220, 232)

DEVKIT = (21, 62, 115)
DEVKIT_FILL = (232, 240, 250)
COMM = (31, 107, 181)
COMM_FILL = (236, 245, 255)
MC2 = (11, 138, 122)
AIN = (91, 75, 213)
HCOMM = (196, 92, 18)
HCOMM_FILL = (255, 246, 236)

ENGINES = {
    "DPU": ((14, 116, 144), "数据处理器", "通信数据面卸载"),
    "AICPU": ((67, 56, 202), "AI CPU 引擎", "Kernel 展开 · TS 调度 · 大带宽"),
    "AIV": ((4, 120, 87), "Vector Core 直驱", "小数据 · 低时延"),
    "CCU": ((180, 83, 9), "集合通信加速单元", "Mission · 微码硬化执行"),
}


def rgb_hex(rgb: tuple[int, int, int]) -> RGBColor:
    return RGBColor(*rgb)


def set_run_font(run, size_pt: float, bold: bool, color: tuple[int, int, int]) -> None:
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = rgb_hex(color)
    run.font.name = LATIN_FONT
    r_pr = run._r.get_or_add_rPr()
    for tag, name in (("a:ea", EA_FONT), ("a:cs", EA_FONT), ("a:latin", LATIN_FONT)):
        node = r_pr.find(qn(tag))
        if node is None:
            node = etree.SubElement(r_pr, qn(tag))
        node.set("typeface", name)


def fill_solid(shape, color: tuple[int, int, int], line=None, line_w_pt: float = 1.25) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb_hex(color)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb_hex(line)
        shape.line.width = Pt(line_w_pt)


def add_round_rect(slide, x, y, w, h, fill, line=None, radius=0.08, line_w=1.25):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    try:
        shape.adjustments[0] = radius
    except Exception:
        pass
    fill_solid(shape, fill, line, line_w)
    return shape


def add_rect(slide, x, y, w, h, fill, line=None, line_w=1.0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    fill_solid(shape, fill, line, line_w)
    return shape


def add_text(
    slide,
    x,
    y,
    w,
    h,
    lines,
    size=12,
    bold=False,
    color=NAVY,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.MIDDLE,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    for i, item in enumerate(lines):
        if isinstance(item, tuple):
            text, item_size, item_bold, item_color = item
        else:
            text, item_size, item_bold, item_color = item, size, bold, color
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = text
        set_run_font(run, item_size, item_bold, item_color)
    return box


def add_arrow(slide, x1, y1, x2, y2, color=HCOMM):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    conn.line.color.rgb = rgb_hex(color)
    conn.line.width = Pt(2.0)
    sp_pr = conn._element.spPr
    ln = sp_pr.find(qn("a:ln"))
    if ln is None:
        ln = etree.SubElement(sp_pr, qn("a:ln"))
    tail = ln.find(qn("a:tailEnd"))
    if tail is None:
        tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return conn


def build_pptx() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, BG)
    add_rect(slide, 0, 0, SLIDE_W, 0.08, RED)
    add_rect(slide, 0, 0.08, 0.12, SLIDE_H - 0.08, DEVKIT)

    add_text(
        slide,
        0.38,
        0.16,
        10.4,
        0.42,
        [("asc-devkit 通信编程能力架构", 26, True, NAVY)],
        anchor=MSO_ANCHOR.BOTTOM,
    )
    add_text(
        slide,
        0.38,
        0.54,
        10.6,
        0.28,
        [("从算子开发套件到多引擎硬件：一层仓、两套编程能力、统一 HCOMM 接口", 12, False, MUTED)],
        anchor=MSO_ANCHOR.TOP,
    )
    add_round_rect(slide, 11.05, 0.22, 1.95, 0.52, RED, radius=0.2)
    add_text(
        slide,
        11.05,
        0.22,
        1.95,
        0.52,
        [("CANN · Ascend C", 11, True, WHITE)],
        align=PP_ALIGN.CENTER,
    )

    # L1 asc-devkit
    add_round_rect(slide, 0.32, 0.92, 12.70, 4.42, DEVKIT_FILL, DEVKIT, 0.04, 1.75)
    add_round_rect(slide, 0.32, 0.92, 12.70, 0.42, DEVKIT, radius=0.04)
    add_rect(slide, 0.32, 1.22, 12.70, 0.12, DEVKIT)
    add_text(
        slide,
        0.48,
        0.92,
        8.6,
        0.42,
        [("asc-devkit    昇腾算子开发套件 · 包含通信仓", 15, True, WHITE)],
    )
    add_text(
        slide,
        9.4,
        0.92,
        3.4,
        0.42,
        [("L1  开发套件", 11, True, WHITE)],
        align=PP_ALIGN.RIGHT,
    )

    # L2 asc-comm
    add_round_rect(slide, 0.50, 1.48, 12.34, 3.70, COMM_FILL, COMM, 0.04, 1.5)
    add_round_rect(slide, 0.50, 1.48, 12.34, 0.38, COMM, radius=0.04)
    add_rect(slide, 0.50, 1.74, 12.34, 0.12, COMM)
    add_text(
        slide,
        0.66,
        1.48,
        8.8,
        0.38,
        [("asc-comm 仓    通信场景开源仓 · 提供 MC² / Ain 编程能力，并接入 HCOMM", 13, True, WHITE)],
    )
    add_text(
        slide,
        9.4,
        1.48,
        3.2,
        0.38,
        [("L2  通信仓", 11, True, WHITE)],
        align=PP_ALIGN.RIGHT,
    )

    # MC2 / Ain
    add_round_rect(slide, 0.72, 2.04, 5.82, 1.42, WHITE, MC2, 0.08, 2.0)
    add_rect(slide, 0.72, 2.12, 0.08, 1.26, MC2)
    add_text(slide, 1.00, 2.10, 5.35, 0.38, [("MC²    通算融合编程能力", 16, True, MC2)])
    add_text(
        slide,
        1.00,
        2.46,
        5.35,
        0.90,
        [
            ("Matrix Computation & Communication", 11, False, MUTED),
            ("计算与通信流水叠加，支撑 AllReduce / AllGather /", 12, False, NAVY),
            ("ReduceScatter 等集合通信与矩阵计算融合算子。", 12, False, NAVY),
        ],
        anchor=MSO_ANCHOR.TOP,
    )

    add_round_rect(slide, 6.74, 2.04, 5.86, 1.42, WHITE, AIN, 0.08, 2.0)
    add_rect(slide, 6.74, 2.12, 0.08, 1.26, AIN)
    add_text(slide, 7.02, 2.10, 5.40, 0.38, [("Ain    单边通信编程能力", 16, True, AIN)])
    add_text(
        slide,
        7.02,
        2.46,
        5.40,
        0.90,
        [
            ("One-sided Put / Get / Signal / Wait", 11, False, MUTED),
            ("对称窗口上发起单边读写与同步，Flush / Wait", 12, False, NAVY),
            ("管理完成语义，面向 MoE / Pipeline 低时延路径。", 12, False, NAVY),
        ],
        anchor=MSO_ANCHOR.TOP,
    )

    add_text(
        slide,
        0.72,
        3.46,
        11.88,
        0.28,
        [("提供编程能力  →  访问统一接口", 11, True, MUTED)],
        align=PP_ALIGN.CENTER,
    )

    # HCOMM
    add_round_rect(slide, 0.72, 3.76, 11.90, 1.22, HCOMM_FILL, HCOMM, 0.08, 2.0)
    add_text(
        slide,
        0.92,
        3.82,
        11.50,
        0.38,
        [("HCOMM    统一编程接口", 16, True, HCOMM)],
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        0.92,
        4.18,
        11.50,
        0.68,
        [
            ("一套 API 屏蔽引擎差异：控制面（通信域 / 通道 / 资源） + 数据面（Write / Read / Notify / Reduce）", 12, False, NAVY),
            ("asc-comm 通过 HCOMM 访问下层通信引擎与硬件加速单元", 12, False, MUTED),
        ],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.TOP,
    )

    add_arrow(slide, 6.66, 5.00, 6.66, 5.46, HCOMM)

    # Engines
    add_text(
        slide,
        0.32,
        5.42,
        12.70,
        0.28,
        [("访问通信引擎  /  硬件加速单元", 12, True, NAVY)],
        align=PP_ALIGN.CENTER,
    )

    names = list(ENGINES.keys())
    gap = 0.22
    box_w = (12.70 - 3 * gap) / 4
    x0 = 0.32
    y0 = 5.72
    for i, name in enumerate(names):
        color, title, desc = ENGINES[name]
        x = x0 + i * (box_w + gap)
        add_round_rect(slide, x, y0, box_w, 1.38, WHITE, color, 0.1, 1.75)
        add_round_rect(slide, x, y0, box_w, 0.10, color, radius=0.02)
        add_text(slide, x, y0 + 0.18, box_w, 0.42, [(name, 20, True, color)], align=PP_ALIGN.CENTER)
        add_text(slide, x + 0.08, y0 + 0.58, box_w - 0.16, 0.28, [(title, 12, True, NAVY)], align=PP_ALIGN.CENTER)
        add_text(slide, x + 0.08, y0 + 0.86, box_w - 0.16, 0.42, [(desc, 11, False, MUTED)], align=PP_ALIGN.CENTER)

    add_rect(slide, 0, 7.28, SLIDE_W, 0.22, NAVY_DEEP)
    add_text(
        slide,
        0.32,
        7.28,
        12.70,
        0.22,
        [("asc-devkit ⊃ asc-comm  →  MC² / Ain 编程  →  HCOMM 统一接口  →  DPU · AICPU · AIV · CCU", 10, False, WHITE)],
        align=PP_ALIGN.CENTER,
    )

    prs.save(PPTX_PATH)


def load_font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_PATH, size)


def to_px(inches: float, axis: str = "x") -> int:
    span = PNG_W if axis == "x" else PNG_H
    total = SLIDE_W if axis == "x" else SLIDE_H
    return int(round(inches / total * span))


def rr(draw: ImageDraw.ImageDraw, box, r, fill, outline=None, width=2):
    draw.rounded_rectangle(box, radius=r, fill=fill, outline=outline, width=width)


def text_center(draw, xy, text, font, fill):
    x0, y0, x1, y1 = xy
    bbox = draw.textbbox((0, 0), text, font=font)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.text((x0 + (x1 - x0 - tw) / 2, y0 + (y1 - y0 - th) / 2 - bbox[1]), text, font=font, fill=fill)


def build_png() -> None:
    img = Image.new("RGB", (PNG_W, PNG_H), BG)
    d = ImageDraw.Draw(img)
    sx = lambda v: to_px(v, "x")
    sy = lambda v: to_px(v, "y")

    font_title = load_font(40)
    font_sub = load_font(20)
    font_h1 = load_font(24)
    font_h2 = load_font(21)
    font_body = load_font(18)
    font_small = load_font(16)
    font_tiny = load_font(15)
    font_eng = load_font(32)
    font_badge = load_font(17)
    font_footer = load_font(15)

    d.rectangle([0, 0, PNG_W, sy(0.08)], fill=RED)
    d.rectangle([0, sy(0.08), sx(0.12), PNG_H], fill=DEVKIT)

    d.text((sx(0.40), sy(0.18)), "asc-devkit 通信编程能力架构", font=font_title, fill=NAVY)
    d.text(
        (sx(0.40), sy(0.56)),
        "从算子开发套件到多引擎硬件：一层仓、两套编程能力、统一 HCOMM 接口",
        font=font_sub,
        fill=MUTED,
    )
    rr(d, [sx(11.05), sy(0.22), sx(13.00), sy(0.74)], 18, RED)
    text_center(d, [sx(11.05), sy(0.22), sx(13.00), sy(0.74)], "CANN · Ascend C", font_badge, WHITE)

    rr(d, [sx(0.32), sy(0.92), sx(13.02), sy(5.34)], 16, DEVKIT_FILL, DEVKIT, 3)
    d.rounded_rectangle([sx(0.32), sy(0.92), sx(13.02), sy(1.34)], radius=16, fill=DEVKIT)
    d.rectangle([sx(0.32), sy(1.18), sx(13.02), sy(1.34)], fill=DEVKIT)
    d.text((sx(0.50), sy(0.99)), "asc-devkit    昇腾算子开发套件 · 包含通信仓", font=font_h1, fill=WHITE)
    layer = "L1  开发套件"
    lb = d.textbbox((0, 0), layer, font=font_small)
    d.text((sx(12.85) - (lb[2] - lb[0]), sy(1.03)), layer, font=font_small, fill=WHITE)

    rr(d, [sx(0.50), sy(1.48), sx(12.84), sy(5.18)], 14, COMM_FILL, COMM, 3)
    d.rounded_rectangle([sx(0.50), sy(1.48), sx(12.84), sy(1.86)], radius=14, fill=COMM)
    d.rectangle([sx(0.50), sy(1.72), sx(12.84), sy(1.86)], fill=COMM)
    d.text(
        (sx(0.68), sy(1.54)),
        "asc-comm 仓    通信场景开源仓 · 提供 MC² / Ain 编程能力，并接入 HCOMM",
        font=font_h2,
        fill=WHITE,
    )
    layer2 = "L2  通信仓"
    lb2 = d.textbbox((0, 0), layer2, font=font_small)
    d.text((sx(12.68) - (lb2[2] - lb2[0]), sy(1.56)), layer2, font=font_small, fill=WHITE)

    rr(d, [sx(0.72), sy(2.04), sx(6.54), sy(3.46)], 14, WHITE, MC2, 4)
    d.rectangle([sx(0.72), sy(2.14), sx(0.80), sy(3.36)], fill=MC2)
    d.text((sx(1.02), sy(2.12)), "MC²    通算融合编程能力", font=font_h1, fill=MC2)
    d.text((sx(1.02), sy(2.50)), "Matrix Computation & Communication", font=font_small, fill=MUTED)
    d.text((sx(1.02), sy(2.78)), "计算与通信流水叠加，支撑 AllReduce / AllGather /", font=font_body, fill=NAVY)
    d.text((sx(1.02), sy(3.06)), "ReduceScatter 等集合通信与矩阵计算融合算子。", font=font_body, fill=NAVY)

    rr(d, [sx(6.74), sy(2.04), sx(12.60), sy(3.46)], 14, WHITE, AIN, 4)
    d.rectangle([sx(6.74), sy(2.14), sx(6.82), sy(3.36)], fill=AIN)
    d.text((sx(7.04), sy(2.12)), "Ain    单边通信编程能力", font=font_h1, fill=AIN)
    d.text((sx(7.04), sy(2.50)), "One-sided Put / Get / Signal / Wait", font=font_small, fill=MUTED)
    d.text((sx(7.04), sy(2.78)), "对称窗口上发起单边读写与同步，Flush / Wait", font=font_body, fill=NAVY)
    d.text((sx(7.04), sy(3.06)), "管理完成语义，面向 MoE / Pipeline 低时延路径。", font=font_body, fill=NAVY)

    mid = "提供编程能力  →  访问统一接口"
    mb = d.textbbox((0, 0), mid, font=font_small)
    d.text(((PNG_W - (mb[2] - mb[0])) / 2, sy(3.52)), mid, font=font_small, fill=MUTED)

    rr(d, [sx(0.72), sy(3.76), sx(12.62), sy(4.98)], 14, HCOMM_FILL, HCOMM, 3)
    text_center(d, [sx(0.72), sy(3.82), sx(12.62), sy(4.20)], "HCOMM    统一编程接口", font_h1, HCOMM)
    t1 = "一套 API 屏蔽引擎差异：控制面（通信域 / 通道 / 资源） + 数据面（Write / Read / Notify / Reduce）"
    t2 = "asc-comm 通过 HCOMM 访问下层通信引擎与硬件加速单元"
    b1 = d.textbbox((0, 0), t1, font=font_body)
    b2 = d.textbbox((0, 0), t2, font=font_body)
    d.text(((PNG_W - (b1[2] - b1[0])) / 2, sy(4.28)), t1, font=font_body, fill=NAVY)
    d.text(((PNG_W - (b2[2] - b2[0])) / 2, sy(4.58)), t2, font=font_tiny, fill=MUTED)

    cx = PNG_W // 2
    y1, y2 = sy(5.00), sy(5.42)
    d.line([(cx, y1), (cx, y2)], fill=HCOMM, width=4)
    d.polygon([(cx - 8, y2 - 12), (cx + 8, y2 - 12), (cx, y2 + 2)], fill=HCOMM)

    cap = "访问通信引擎  /  硬件加速单元"
    cb = d.textbbox((0, 0), cap, font=font_small)
    d.text(((PNG_W - (cb[2] - cb[0])) / 2, sy(5.44)), cap, font=font_small, fill=NAVY)

    names = list(ENGINES.keys())
    gap = sx(0.22)
    box_w = (sx(12.70) - 3 * gap) / 4
    x0 = sx(0.32)
    y0 = sy(5.72)
    h = sy(1.38)
    for i, name in enumerate(names):
        color, title, desc = ENGINES[name]
        x = x0 + i * (box_w + gap)
        rr(d, [x, y0, x + box_w, y0 + h], 16, WHITE, color, 3)
        d.rounded_rectangle([x, y0, x + box_w, y0 + sy(0.10)], radius=8, fill=color)
        text_center(d, [x, y0 + sy(0.16), x + box_w, y0 + sy(0.58)], name, font_eng, color)
        text_center(d, [x, y0 + sy(0.58), x + box_w, y0 + sy(0.88)], title, font_body, NAVY)
        text_center(d, [x, y0 + sy(0.88), x + box_w, y0 + sy(1.26)], desc, font_tiny, MUTED)

    d.rectangle([0, sy(7.28), PNG_W, PNG_H], fill=NAVY_DEEP)
    foot = "asc-devkit ⊃ asc-comm  →  MC² / Ain 编程  →  HCOMM 统一接口  →  DPU · AICPU · AIV · CCU"
    fb = d.textbbox((0, 0), foot, font=font_footer)
    d.text(((PNG_W - (fb[2] - fb[0])) / 2, sy(7.33)), foot, font=font_footer, fill=WHITE)

    img.save(PNG_PATH, "PNG")


if __name__ == "__main__":
    build_pptx()
    build_png()
    print(f"wrote {PPTX_PATH}")
    print(f"wrote {PNG_PATH}")
